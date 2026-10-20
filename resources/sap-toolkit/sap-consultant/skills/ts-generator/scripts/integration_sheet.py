#!/usr/bin/env python3
"""Write the integration team's "Entegrasyon Teknik Detay" sheet from the TS interview.

    py integration_sheet.py <answers.json> <out.xlsx> [--lang tr|en]

The PO/CPI team works from a one-sheet form (service, trigger, volume, frequency, timeout,
source and target with their connection methods, sync/async, request and response samples,
the other system's documentation). ts-generator asks those questions in the same round as
the TS and writes this file beside it (references/INTEGRATION_SHEET.md says when and how).

Their form, kept in its order, with three additions that do not change what they fill in:
  - a Source column per row (FS / TS / User / Proposal / Open), so a reviewer sees what was
    derived and what was confirmed, and nothing derived passes for an answer;
  - an "Additional Information" block (authentication, error handling and retry, monitoring,
    endpoints per environment, personal data) -- the questions their review asks anyway;
  - a Field Mapping sheet, from the TS's section 2.9.
The service and integration ids are optional: not every project keeps them in Trace.

answers.json:
  {"wricef": "...", "title": "...", "status": "Taslak", "date": "2026-10-09",
   "fields": {"<key>": {"value": "...", "source": "fs|ts|user|proposal|open"}, ...},
   "mapping": [{"source": "...", "target": "...", "type": "...", "required": true, "rule": "..."}]}
Keys: see FIELDS below. A field left out or empty is written as Open.

The words are the master's (plugins/office-tools/lib/doc_language.json, table
"integration"); scripts/test_doc_language.py holds WORDS equal to it. openpyxl only.
"""
from __future__ import annotations

import datetime as _dt
import json
import sys
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

# Turkish -> English, the master's "integration" table. Change the words THERE.
WORDS = {
    "Entegrasyon Teknik Detay": "Integration Technical Detail",
    "Entegrasyon Genel Bilgiler": "Integration General Information",
    "Ek Bilgiler": "Additional Information",
    "Alan Eşleme": "Field Mapping",
    "Alan": "Field", "Değer": "Value", "Kaynak": "Source",
    "Başlık": "Title", "Durum": "Status", "Tarih": "Date", "Taslak": "Draft",
    "Servis ID": "Service ID", "Entegrasyon ID": "Integration ID", "Uygulama Linki": "Application Link",
    "Öngörülen Kullanım Hacmi": "Expected Volume", "Çalışma Şekli": "Trigger Mode",
    "Çalışma Sıklığı": "Frequency", "Timeout Süresi": "Timeout",
    "Kaynak Sistem": "Source System", "Kaynak Sistem Bağlantı Yöntemi": "Source System Connection Method",
    "Hedef Sistem": "Target System", "Hedef Sistem Bağlantı Yöntemi": "Target System Connection Method",
    "Senkron/Asenkron": "Synchronous/Asynchronous",
    "Request Örnek Veri Deseni": "Request Sample Payload",
    "Response Örnek Veri Deseni": "Response Sample Payload",
    "Entegre olunan sistemden sağlanan doküman bilgisi": "Documentation provided by the integrated system",
    "Kimlik Doğrulama": "Authentication", "Hata Yönetimi ve Yeniden Deneme": "Error Handling and Retry",
    "İzleme ve Alarm": "Monitoring and Alerting",
    "Ortam Uç Noktaları (DEV/QA/PRD)": "Endpoints per Environment (DEV/QA/PRD)",
    "Kişisel Veri (KVKK)": "Personal Data (KVKK/GDPR)",
    "Kaynak Alan": "Source Field", "Hedef Alan": "Target Field", "Tip": "Type", "Zorunlu": "Required",
    "Dönüşüm Kuralı": "Transformation Rule", "Evet": "Yes", "Hayır": "No",
    "FS": "FS", "TS": "TS",
    "Kullanıcı": "User", "Öneri": "Proposal", "Açık": "Open",
    "* zorunlu alan": "* required field",
    "Kaynak sütunu: FS = fonksiyonel spesifikasyondan, TS = teknik spesifikasyondan, Kullanıcı = görüşmede verildi, Öneri = türetildi, onay bekliyor, Açık = henüz cevap yok.":
        "Source column: FS = from the functional specification, TS = from the technical specification, User = given in the interview, Proposal = derived, awaiting confirmation, Open = no answer yet.",
}

# (key, Turkish label, required) in the form's own order.
FIELDS = [
    ("service_id", "Servis ID", False),
    ("integration_id", "Entegrasyon ID", False),
    ("app_link", "Uygulama Linki", False),
    ("volume", "Öngörülen Kullanım Hacmi", True),
    ("trigger", "Çalışma Şekli", True),
    ("frequency", "Çalışma Sıklığı", True),
    ("timeout", "Timeout Süresi", True),
    ("source_system", "Kaynak Sistem", True),
    ("source_method", "Kaynak Sistem Bağlantı Yöntemi", True),
    ("target_system", "Hedef Sistem", True),
    ("target_method", "Hedef Sistem Bağlantı Yöntemi", True),
    ("sync", "Senkron/Asenkron", True),
    ("request_sample", "Request Örnek Veri Deseni", True),
    ("response_sample", "Response Örnek Veri Deseni", True),
    ("partner_docs", "Entegre olunan sistemden sağlanan doküman bilgisi", False),
]
EXTRA = [
    ("auth", "Kimlik Doğrulama"),
    ("error_handling", "Hata Yönetimi ve Yeniden Deneme"),
    ("monitoring", "İzleme ve Alarm"),
    ("endpoints", "Ortam Uç Noktaları (DEV/QA/PRD)"),
    ("personal_data", "Kişisel Veri (KVKK)"),
]
SOURCES = {"fs": "FS", "ts": "TS", "user": "Kullanıcı", "proposal": "Öneri", "open": "Açık"}

ACCENT, DEEP, INK, MUTED, RULE, TINT = "0072BC", "00395E", "22262A", "626C76", "D8E3EC", "EAF4FB"
OPEN_FILL, PROPOSAL_FILL = "FDECEA", "FFF8E1"
FONT = "Segoe UI"


def take_lang(argv: list[str]) -> str:
    """--lang tr|en, taken out of argv (default tr)."""
    lang = "tr"
    if "--lang" in argv:
        i = argv.index("--lang")
        if i + 1 >= len(argv) or argv[i + 1] not in ("tr", "en"):
            raise SystemExit("--lang takes tr or en")
        lang = argv[i + 1]
        del argv[i:i + 2]
    return lang


def word(tr: str, lang: str) -> str:
    return tr if lang == "tr" else WORDS[tr]


def build(data: dict, out: Path, lang: str = "tr") -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    w = lambda tr: word(tr, lang)  # noqa: E731
    hair = Side(style="thin", color=RULE)
    border = Border(top=hair, bottom=hair, left=hair, right=hair)
    top = Alignment(wrap_text=True, vertical="top", indent=1)
    mono = Alignment(wrap_text=True, vertical="top")

    def font(size=10, bold=False, color=INK, italic=False, name=FONT):
        return Font(name=name, size=size, bold=bold, color=color, italic=italic)

    wb = Workbook()
    ws = wb.active
    ws.title = w("Entegrasyon Teknik Detay")[:31]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 78
    ws.column_dimensions["C"].width = 14

    ws["A1"] = w("Entegrasyon Teknik Detay")
    ws["A1"].font = font(16, True, DEEP)
    ws.row_dimensions[1].height = 26
    meta = [("WRICEF ID", data.get("wricef", "")), (w("Başlık"), data.get("title", "")),
            (w("Durum"), data.get("status") or w("Taslak")),
            (w("Tarih"), data.get("date") or _dt.date.today().isoformat())]
    r = 2
    for k, v in meta:
        ws.cell(r, 1, k).font = font(9, True, MUTED)
        ws.cell(r, 2, str(v)).font = font(10)
        r += 1
    r += 1

    fields = data.get("fields") or {}
    counts = {"open": 0, "proposal": 0, "filled": 0}

    def header(row: int, text: str) -> None:
        for c, t in enumerate((text, w("Değer"), w("Kaynak")), 1):
            cell = ws.cell(row, c, t)
            cell.font = font(10, True, "FFFFFF")
            cell.fill = PatternFill("solid", fgColor=ACCENT)
            cell.alignment = Alignment(vertical="center", indent=1)
            cell.border = border
        ws.row_dimensions[row].height = 20

    def line(row: int, key: str, label: str, required: bool) -> None:
        f = fields.get(key) or {}
        if isinstance(f, str):
            f = {"value": f}
        value = str(f.get("value") or "").strip()
        src = str(f.get("source") or ("open" if not value else "user")).lower()
        if src not in SOURCES:
            src = "user"
        if not value:
            src = "open"
        counts["open" if src == "open" else "proposal" if src == "proposal" else "filled"] += 1
        a = ws.cell(row, 1, label + ("*" if required else ""))
        a.font = font(10, True)
        b = ws.cell(row, 2, value if value else w("Açık"))
        b.font = font(9, name="Consolas") if key.endswith("_sample") and value else font(10, italic=not value, color=INK if value else MUTED)
        c = ws.cell(row, 3, w(SOURCES[src]))
        c.font = font(9, color=MUTED)
        fill = OPEN_FILL if src == "open" else PROPOSAL_FILL if src == "proposal" else None
        for cell in (a, b, c):
            cell.alignment = mono if cell is b and key.endswith("_sample") else top
            cell.border = border
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
        # The label wraps too (column A holds ~34 characters); a cut label reads as a missing one.
        lines = max(1, value.count("\n") + 1, len(value) // 70 + 1, len(label) // 34 + 1)
        ws.row_dimensions[row].height = min(400, (13 if key.endswith("_sample") and value else 15) * lines + 6)

    header(r, w("Entegrasyon Genel Bilgiler")); r += 1
    for key, label, req in FIELDS:
        line(r, key, w(label), req); r += 1
    r += 1
    header(r, w("Ek Bilgiler")); r += 1
    for key, label in EXTRA:
        line(r, key, w(label), False); r += 1
    r += 1
    ws.cell(r, 1, w("* zorunlu alan")).font = font(9, italic=True, color=MUTED); r += 1
    note = ws.cell(r, 1, w("Kaynak sütunu: FS = fonksiyonel spesifikasyondan, TS = teknik spesifikasyondan, Kullanıcı = görüşmede verildi, Öneri = türetildi, onay bekliyor, Açık = henüz cevap yok."))
    note.font = font(9, italic=True, color=MUTED)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    m = wb.create_sheet(w("Alan Eşleme")[:31])
    m.sheet_view.showGridLines = False
    cols = [("Kaynak Alan", 30), ("Hedef Alan", 30), ("Tip", 16), ("Zorunlu", 11), ("Dönüşüm Kuralı", 60)]
    for i, (t, width) in enumerate(cols, 1):
        cell = m.cell(1, i, w(t))
        cell.font = font(10, True, "FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.border = border
        cell.alignment = Alignment(vertical="center", indent=1)
        m.column_dimensions[cell.column_letter].width = width
    rows = data.get("mapping") or []
    for j, row in enumerate(rows, 2):
        req = row.get("required")
        vals = [row.get("source", ""), row.get("target", ""), row.get("type", ""),
                "" if req is None else w("Evet") if req else w("Hayır"), row.get("rule", "")]
        for i, v in enumerate(vals, 1):
            cell = m.cell(j, i, str(v))
            cell.font = font(10)
            cell.alignment = top
            cell.border = border
            if j % 2:
                cell.fill = PatternFill("solid", fgColor=TINT)
    if not rows:
        m.cell(2, 1, w("Açık")).font = font(10, italic=True, color=MUTED)
    m.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"open": counts["open"], "proposal": counts["proposal"], "filled": counts["filled"], "mapping": len(rows)}


def main(argv: list[str]) -> int:
    argv = list(argv)
    lang = take_lang(argv)
    if len(argv) != 3:
        print(__doc__.split("\n\n")[0].strip())
        return 2
    data = json.loads(Path(argv[1]).read_text(encoding="utf-8-sig"))
    out = Path(argv[2])
    res = build(data, out, lang)
    print(f"written: {out}")
    print(f"fields filled {res['filled']}, proposals {res['proposal']}, open {res['open']}; mapping rows {res['mapping']}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
