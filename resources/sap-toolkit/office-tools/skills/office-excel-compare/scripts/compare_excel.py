#!/usr/bin/env python3
"""Compare two Excel/CSV files and produce a colour-coded delta report."""

import argparse
import sys
from pathlib import Path


def load(path: Path, sheet, parse_dates, engine: str):
    suffix = path.suffix.lower()
    if engine == "polars":
        import polars as pl
        if suffix in (".csv", ".tsv"):
            sep = "\t" if suffix == ".tsv" else ","
            return pl.read_csv(path, separator=sep)
        return pl.read_excel(path, sheet_name=sheet or 0)
    else:
        import pandas as pd
        dates = [c.strip() for c in parse_dates.split(",")] if parse_dates else False
        if suffix in (".csv", ".tsv", ".txt"):
            sep = "\t" if suffix in (".tsv", ".txt") else ","
            return pd.read_csv(path, sep=sep, parse_dates=dates or False)
        return pd.read_excel(path, sheet_name=sheet or 0,
                             parse_dates=dates or False, engine="openpyxl")


def compare_pandas(df1, df2, key: str, normalize: bool):
    import pandas as pd

    # normalize keys
    if normalize:
        df1[key] = df1[key].astype(str).str.strip().str.lower()
        df2[key] = df2[key].astype(str).str.strip().str.lower()

    # check for duplicate keys
    dup1 = df1[df1.duplicated(subset=[key], keep=False)]
    dup2 = df2[df2.duplicated(subset=[key], keep=False)]

    # common columns only for value comparison
    cols1 = set(df1.columns) - {key}
    cols2 = set(df2.columns) - {key}
    added_cols = cols2 - cols1
    removed_cols = cols1 - cols2
    common_cols = cols1 & cols2

    keys1 = set(df1[key].astype(str))
    keys2 = set(df2[key].astype(str))

    added_keys = keys2 - keys1
    removed_keys = keys1 - keys2
    common_keys = keys1 & keys2

    # rows added in file2
    added_rows = df2[df2[key].astype(str).isin(added_keys)].copy()
    added_rows["_status"] = "ADDED"

    # rows removed from file1
    removed_rows = df1[df1[key].astype(str).isin(removed_keys)].copy()
    removed_rows["_status"] = "REMOVED"

    # changed rows — compare common keys on common columns
    changed_rows = []
    changed_detail = []
    if common_cols:
        m1 = df1[df1[key].astype(str).isin(common_keys)].set_index(key)[list(common_cols)]
        m2 = df2[df2[key].astype(str).isin(common_keys)].set_index(key)[list(common_cols)]
        # align
        m1, m2 = m1.align(m2, join="inner")
        diff_mask = (m1 != m2) & ~(m1.isna() & m2.isna())
        changed_idx = diff_mask.any(axis=1)
        for idx in m1[changed_idx].index:
            row = {key: idx, "_status": "CHANGED"}
            for col in common_cols:
                old_val = m1.loc[idx, col]
                new_val = m2.loc[idx, col]
                if old_val != new_val and not (pd.isna(old_val) and pd.isna(new_val)):
                    row[f"{col}_OLD"] = old_val
                    row[f"{col}_NEW"] = new_val
            changed_rows.append(row)

    changed_df = pd.DataFrame(changed_rows) if changed_rows else pd.DataFrame()

    summary = {
        "file1_rows": len(df1),
        "file2_rows": len(df2),
        "added_rows": len(added_rows),
        "removed_rows": len(removed_rows),
        "changed_rows": len(changed_df),
        "unchanged_rows": len(common_keys) - len(changed_df),
        "added_columns": sorted(added_cols),
        "removed_columns": sorted(removed_cols),
        "duplicate_keys_file1": len(dup1),
        "duplicate_keys_file2": len(dup2),
    }

    return summary, added_rows, removed_rows, changed_df, dup1, dup2


def write_delta_report(out_path: Path, summary: dict, added, removed, changed, key: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, PatternFill

    wb = Workbook()

    # --- Summary sheet ---
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.column_dimensions["A"].width = 30
    ws_sum.column_dimensions["B"].width = 20
    title_font = Font(bold=True, size=13)
    ws_sum["A1"] = "Delta Report Summary"
    ws_sum["A1"].font = title_font
    ws_sum.merge_cells("A1:B1")
    rows = [
        ("File 1 rows", summary["file1_rows"]),
        ("File 2 rows", summary["file2_rows"]),
        ("", ""),
        ("Added rows (in File 2, not File 1)", summary["added_rows"]),
        ("Removed rows (in File 1, not File 2)", summary["removed_rows"]),
        ("Changed rows (same key, different values)", summary["changed_rows"]),
        ("Unchanged rows", summary["unchanged_rows"]),
        ("", ""),
        ("Added columns", ", ".join(summary["added_columns"]) or "none"),
        ("Removed columns", ", ".join(summary["removed_columns"]) or "none"),
        ("Duplicate keys in File 1", summary["duplicate_keys_file1"]),
        ("Duplicate keys in File 2", summary["duplicate_keys_file2"]),
    ]
    fills = {
        "Added rows (in File 2, not File 1)": "C6EFCE",
        "Removed rows (in File 1, not File 2)": "FFC7CE",
        "Changed rows (same key, different values)": "FFEB9C",
    }
    for i, (label, value) in enumerate(rows, start=3):
        ws_sum.cell(row=i, column=1, value=label)
        ws_sum.cell(row=i, column=2, value=value)
        if label in fills:
            fill = PatternFill("solid", fgColor=fills[label])
            ws_sum.cell(row=i, column=1).fill = fill
            ws_sum.cell(row=i, column=2).fill = fill

    def _write_sheet(ws, df, status_color: str, title: str):
        if df is None or len(df) == 0:
            ws.cell(row=1, column=1, value=f"No {title.lower()}")
            return
        fill = PatternFill("solid", fgColor=status_color)
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(bold=True, color="FFFFFF")
        cols = list(df.columns)
        for ci, col in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=col)
            cell.fill = header_fill
            cell.font = header_font
        for ri, row in enumerate(df.itertuples(index=False), 2):
            for ci, val in enumerate(row, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.fill = fill
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    import pandas as pd
    ws_add = wb.create_sheet("Added")
    _write_sheet(ws_add, added if len(added) > 0 else pd.DataFrame(), "C6EFCE", "Added")

    ws_rem = wb.create_sheet("Removed")
    _write_sheet(ws_rem, removed if len(removed) > 0 else pd.DataFrame(), "FFC7CE", "Removed")

    ws_chg = wb.create_sheet("Changed")
    _write_sheet(ws_chg, changed if len(changed) > 0 else pd.DataFrame(), "FFEB9C", "Changed")

    wb.save(out_path)


def print_summary(summary: dict, path1, path2):
    print(f"\n{'─'*55}")
    print(f"  Delta Summary")
    print(f"{'─'*55}")
    print(f"  File 1 : {path1}  ({summary['file1_rows']:,} rows)")
    print(f"  File 2 : {path2}  ({summary['file2_rows']:,} rows)")
    print(f"{'─'*55}")
    print(f"  Added   (in File 2, not File 1) : {summary['added_rows']:>6,}")
    print(f"  Removed (in File 1, not File 2) : {summary['removed_rows']:>6,}")
    print(f"  Changed (same key, diff values) : {summary['changed_rows']:>6,}")
    print(f"  Unchanged                       : {summary['unchanged_rows']:>6,}")
    if summary["added_columns"]:
        print(f"\n  Columns added  : {summary['added_columns']}")
    if summary["removed_columns"]:
        print(f"  Columns removed: {summary['removed_columns']}")
    if summary["duplicate_keys_file1"] or summary["duplicate_keys_file2"]:
        print(f"\n  ⚠ Duplicate keys — File 1: {summary['duplicate_keys_file1']}, "
              f"File 2: {summary['duplicate_keys_file2']}")
        print(f"    Key column may not uniquely identify rows — results may be inaccurate")
    print(f"{'─'*55}\n")


def main():
    ap = argparse.ArgumentParser(description="Compare two Excel/CSV files and produce a delta report")
    ap.add_argument("--file1", required=True, help="Before file (.xlsx or .csv)")
    ap.add_argument("--file2", required=True, help="After file (.xlsx or .csv)")
    ap.add_argument("--key", required=True, help="Column that uniquely identifies each row")
    ap.add_argument("--sheet1", default=None, help="Sheet in file1 (default: first)")
    ap.add_argument("--sheet2", default=None, help="Sheet in file2 (default: first)")
    ap.add_argument("--output", "-o", default="delta_report.xlsx")
    ap.add_argument("--summary-only", action="store_true",
                    help="Print summary only — do not write row data to output")
    ap.add_argument("--normalize-keys", action="store_true",
                    help="Strip whitespace and lowercase keys before comparing")
    ap.add_argument("--parse-dates", default=None)
    ap.add_argument("--engine", choices=["pandas", "polars"], default="pandas")
    args = ap.parse_args()

    p1 = Path(args.file1)
    p2 = Path(args.file2)
    for p in (p1, p2):
        if not p.exists():
            print(f"ERROR: file not found: {p}", file=sys.stderr)
            sys.exit(1)

    df1 = load(p1, args.sheet1, args.parse_dates, args.engine)
    df2 = load(p2, args.sheet2, args.parse_dates, args.engine)

    # convert polars to pandas for comparison logic
    if args.engine == "polars":
        df1 = df1.to_pandas()
        df2 = df2.to_pandas()

    if args.key not in df1.columns:
        print(f"ERROR: key column '{args.key}' not found in File 1.\n"
              f"Available: {list(df1.columns)}", file=sys.stderr)
        sys.exit(1)
    if args.key not in df2.columns:
        print(f"ERROR: key column '{args.key}' not found in File 2.\n"
              f"Available: {list(df2.columns)}", file=sys.stderr)
        sys.exit(1)

    summary, added, removed, changed, dup1, dup2 = compare_pandas(
        df1, df2, args.key, args.normalize_keys
    )

    print_summary(summary, p1.name, p2.name)

    if not args.summary_only:
        out_path = Path(args.output)
        write_delta_report(out_path, summary, added, removed, changed, args.key)
        size_kb = out_path.stat().st_size / 1024
        print(f"✓ Delta report: {out_path}  ({size_kb:.1f} KB)")
        print(f"  Sheets: Summary | Added (green) | Removed (red) | Changed (yellow)")


if __name__ == "__main__":
    main()
