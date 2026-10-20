"""office-docx — Markdown (subset) -> native, editable Word .docx via python-docx.

Mirrors the office-pdf input (Markdown) but produces an editable .docx carrying
the NTT DATA house theme: the mark repeating in the header of every page, one
accent running through headings, rules, table headers and links, and typography
set for reading rather than for fitting. Supported Markdown subset:
  # ... ######        headings (Word Heading 1-6)
  - / * bullet         (nesting by 2-space indent)
  1. ordered list
  | a | b |            GFM tables (the row of --- separators is the divider)
  ```code```           fenced code block (monospace, shaded)
  ```mermaid           the diagram as a picture (Chromium); code if none
  > quote
  ---                  horizontal rule (a hairline, as in the PDF)
  **bold**  *italic*  `code`   inline runs
  [label](url)  <url>  url     clickable hyperlinks
  ![alt](path)         image on its own line (path relative to the .md,
                       centred, shrunk to the text column if oversized)
  blank line           paragraph break

Usage:
    py build_docx.py --input report.md --output report.docx --title "Report"
    py build_docx.py --input r.md --output r.docx --redact-pii
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

# A Turkish output path (ı, ş, ğ) cannot be encoded by the default Windows cp1252
# console, and the only thing this script prints is that path -- so the document
# was written and the process still died on the success line. Caught by the
# showcase render, which builds ~40 Turkish-named files through a pipe and read
# 29 of them as failures. Same shim as office-excel-images.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "..", "lib"))
try:
    from redact import redact_text, count_matches
except Exception:
    redact_text = None
    count_matches = None

# The palette is shared with every other Office renderer -- see lib/theme.py for
# why it is not defined here. Not optional the way redact.py is: a missing theme
# means a document in the wrong brand, which is worse than a loud failure.
from theme import (ACCENT, ACCENT_HEX, BAND_HEX, BODY_FONT, DEEP,  # noqa: E402
                   INK, LIGHT_FONT, MONO_FONT, MUTED, RULE_HEX, SEMI_FONT,
                   darken, logo_path, parse_hex)

# ---------------------------------------------------------- active palette ---
# The house accent and heading face, unless the caller overrides them. Module
# state rather than a threaded parameter because `_add_hyperlink` needs the
# accent and sits four calls below `build()`, and because this script is one
# build per process: `main()` runs once, and render_showcase.py invokes it as a
# SUBPROCESS per page rather than importing it.
#
# The override exists for a surface that already has a designed identity of its
# own. The showcase pages ship as .html and .docx side by side and are set in
# Cambria over a warm paper with a restrained green -- a written decision, not a
# default nobody revisited. Forcing NTT blue and Segoe UI on the Word half would
# have split one page into two looks, which is the failure the pairing exists to
# prevent. So the surface keeps its palette and gains the machinery: the mark in
# the header, real page fields, hairline tables, callouts, the Turkish font slots.
_ACCENT = ACCENT
_ACCENT_HEX = ACCENT_HEX
_DEEP = DEEP
_HEAD_FONT = SEMI_FONT


def _apply_palette(accent: str = "", heading_font: str = "") -> None:
    """Point the module at this build's palette. Empty argument = house default."""
    global _ACCENT, _ACCENT_HEX, _DEEP, _HEAD_FONT
    if accent:
        _ACCENT = parse_hex(accent)
        _ACCENT_HEX = "%02X%02X%02X" % _ACCENT
        _DEEP = darken(_ACCENT)
    else:
        _ACCENT, _ACCENT_HEX, _DEEP = ACCENT, ACCENT_HEX, DEEP
    _HEAD_FONT = heading_font or SEMI_FONT

# Written into the .docx core properties so a consumer can ask "is this still
# the markdown it was built from?" without rebuilding and diffing binaries.
SOURCE_STAMP = "office-docx source-sha:"


def source_sha(md_text: str) -> str:
    """Hash of the DECODED markdown, so CRLF/LF never shows up as a difference."""
    import hashlib
    return hashlib.sha256(md_text.replace("\r\n", "\n").encode("utf-8")).hexdigest()[:16]


def stamped_sha(docx_path) -> str | None:
    """The source hash a .docx was stamped with, or None if it carries no stamp."""
    try:
        from docx import Document
        note = Document(str(docx_path)).core_properties.comments or ""
    except Exception:
        return None
    _, sep, sha = note.partition(SOURCE_STAMP)
    return sha.strip() or None if sep else None
_LINK = r"\[[^\]\n]+\]\([^)\s]+\)"      # [label](url)
_AUTOLINK = r"<https?://[^>\s]+>"        # <url>
_BARE_URL = r"https?://[^\s<>)\]]+"      # url on its own
_BOLD = r"\*\*(?:[^*]|\*(?!\*))+?\*\*"   # **bold**, may contain a lone *
_ITALIC = r"\*(?!\*)(?:[^*]|\*\*)+?\*(?!\*)"   # *italic*, may contain **bold**
_INLINE = re.compile(
    rf"({_LINK}|{_AUTOLINK}|{_BOLD}|{_ITALIC}|`.+?`|{_BARE_URL})"
)


def _require_docx():
    try:
        import docx  # noqa: F401
    except ImportError:
        sys.exit("ERROR: python-docx required. Run: py -m pip install python-docx")


def _face(target, name, size=None, bold=None, color=None):
    """Set a font on ALL FOUR rFonts slots of a run or a style.

    python-docx writes only `w:ascii`. Turkish is full of characters outside
    Latin-1 -- the dotless i above all -- and for those Word consults `w:hAnsi`
    and `w:eastAsia` instead, silently substituting a different face. The result
    is a single Turkish word set in two fonts, which is exactly the kind of
    defect nobody reports and everybody notices.
    """
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    target.font.name = name
    rfonts = target._element.rPr.rFonts
    for slot in ("w:eastAsia", "w:hAnsi", "w:cs"):
        rfonts.set(qn(slot), name)
    if size is not None:
        target.font.size = Pt(size)
    if bold is not None:
        target.font.bold = bold
    if color is not None:
        target.font.color.rgb = RGBColor(*color)
    return target


def _shade(target, hex_fill):
    """Background fill for a table cell or a paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    props = (target._tc.get_or_add_tcPr() if hasattr(target, "_tc")
             else target._p.get_or_add_pPr())
    props.append(el)


def _rule(paragraph, edge="bottom", hex_color=RULE_HEX, size=6, space=6):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    borders = pPr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        pPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), hex_color)
    borders.append(el)


def _field(paragraph, instr):
    """A Word field (PAGE / NUMPAGES) — python-docx has no API for these.

    Real fields, not baked numbers: the page count has to survive the reader
    editing the document, which is the whole point of shipping .docx.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    return run


def _add_hyperlink(paragraph, url, label):
    """Insert a real, clickable Word hyperlink — python-docx has no API for it,
    so the w:hyperlink element and its relationship are built by hand."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _ACCENT_HEX)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)

    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _add_inline(paragraph, text, bold=False, italic=False):
    """Render **bold**, *italic*, `code` and links as runs in a paragraph.

    Recursive, so markers nest: `**\\`code\\`**` is bold monospace rather than a
    bold run with literal backticks in it. A code span is a leaf — markers inside
    it are content, not formatting."""
    for tok in _INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("[") and tok.endswith(")") and "](" in tok:
            label, _, url = tok[1:-1].partition("](")
            _add_hyperlink(paragraph, _unpark(url), _unpark(label))
        elif tok.startswith("<") and tok.endswith(">"):
            _add_hyperlink(paragraph, _unpark(tok[1:-1]), _unpark(tok[1:-1]))
        elif tok.startswith(("http://", "https://")):
            trail = ""
            while tok and tok[-1] in ".,;:!?":
                trail, tok = tok[-1] + trail, tok[:-1]
            _add_hyperlink(paragraph, _unpark(tok), _unpark(tok))
            if trail:
                _add_inline(paragraph, trail, bold, italic)
        elif len(tok) > 4 and tok.startswith("**") and tok.endswith("**"):
            _add_inline(paragraph, tok[2:-2], True, italic)
        elif len(tok) > 2 and tok.startswith("*") and tok.endswith("*"):
            _add_inline(paragraph, tok[1:-1], bold, True)
        elif len(tok) > 2 and tok.startswith("`") and tok.endswith("`"):
            r = paragraph.add_run(_unpark(tok[1:-1]))
            r.font.name = "Consolas"
            r.bold, r.italic = bold or None, italic or None
        else:
            r = paragraph.add_run(_unpark(tok))
            r.bold, r.italic = bold or None, italic or None


def _fresh_num_id(doc, style_name="List Number"):
    """Clone the list style's numbering into a new instance.

    Every "List Number" paragraph otherwise shares one numbering instance, so
    Word numbers them continuously across the whole document — three separate
    3-step lists come out as 1-3, 4-6, 7-9. A fresh numId per list restarts at 1.
    Returns None if the template has no numbering part (then Word's default
    behaviour stands)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    try:
        numbering = doc.part.numbering_part.element
        style = doc.styles[style_name].element
        path = f'{qn("w:pPr")}/{qn("w:numPr")}/{qn("w:numId")}'
        base_id = style.find(path).get(qn("w:val"))
        abstract = None
        used = []
        for num in numbering.findall(qn("w:num")):
            used.append(int(num.get(qn("w:numId"))))
            if num.get(qn("w:numId")) == base_id:
                abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        if abstract is None:
            return None
        new_id = max(used) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abstract)
        num.append(a)
        numbering.append(num)
        return new_id
    except Exception:
        return None


def _apply_num(paragraph, num_id, level=0):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if num_id is None:
        return
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numPr.append(ilvl)
    numPr.append(nid)
    paragraph._p.get_or_add_pPr().append(numPr)


def _is_table_sep(line):
    return bool(re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", line)) and "-" in line


# Markdown's backslash escapes: `\|` inside a table cell, `\*` for a literal star.
# They were printed with the backslash and a `\|` split its cell in two -- the code
# quality board escapes every pipe in its findings, and Kokpit's documents escape
# freely (2026-09-29). Outside code, each escape is parked on a private-use
# character while the line is parsed and turned back into the plain character
# when the run is written.
_ESCAPABLE = "\\*_`[]#|<>"
_PARK = {c: chr(0xE000 + i) for i, c in enumerate(_ESCAPABLE)}
_UNPARK = {v: k for k, v in _PARK.items()}
_ESC_RE = re.compile(r"\\([%s])" % re.escape(_ESCAPABLE))


def _park(line: str) -> str:
    return _ESC_RE.sub(lambda m: _PARK[m.group(1)], line)


def _unpark(text: str) -> str:
    return "".join(_UNPARK.get(ch, ch) for ch in text) if text else text


def _park_outside_fences(lines: list[str]) -> list[str]:
    out, fence = [], False
    for ln in lines:
        if ln.strip().startswith("```"):
            fence = not fence
            out.append(ln)
        else:
            out.append(ln if fence else _park(ln))
    return out


def _split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


_HTML_COMMENT = re.compile(r"<!--.*?-->", re.S)


def strip_html_comments(md_text: str) -> str:
    """Drop `<!-- ... -->` blocks. Every markdown renderer does; this one did not.

    They are how a source file carries instructions to whoever rebuilds it -- the
    global manual's own header holds its rebuild command that way. Rendering them
    puts that command in the middle of a document a consultant reads.

    NOT applied to source_sha(): the stamp hashes the file as written, so adding
    or editing a comment still counts as the source having changed. That is the
    honest reading -- the note is part of the file even when it is not part of the
    document.
    """
    return _HTML_COMMENT.sub("", md_text)


# `#` to `######`. Only three were recognised until 2026-09-28, and a multi-chapter
# BBP puts its sub-sections (`2.1 Current Situation`) at `####` -- they came out
# of Word as body paragraphs with the hashes showing, invisible to the navigation
# pane and to any table of contents.
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
# The lines that end a running paragraph: a heading, a list item, a quote, a fence
# or a rule. Kept beside HEADING so the two cannot disagree about depth again.
_BLOCK_START = re.compile(r"^(#{1,6}\s|[-*]\s|\d+\.\s|>|```|---)")


def _plain(text: str) -> str:
    """A heading or title as a reader sees it: no markup, one space, any case."""
    return " ".join(re.sub(r"[*_`]", "", text).split()).casefold()


def take_leading_title(lines: list[str], title: str) -> tuple[list[str], str]:
    """With --title, the title block owns the document title.

    A skill's Markdown opens with `# <title>` because office-pdf and a reader of
    the .md need one, and the same file then goes through here with --title --
    so Word printed the title twice, once as the title block and once as a
    Heading 1 under it (BBP sample, 2026-09-28). One owner: when --title is
    given, a `#` heading that is the FIRST content line is the document's title
    and is not rendered again. Returns the remaining lines and the text of the
    heading taken ("" when none was).

    Only a single `#` and only in first position: a `## 1. Purpose` or a `#`
    further down is structure, not the title.
    """
    if not title:
        return lines, ""
    for idx, line in enumerate(lines):
        if not line.strip():
            continue
        m = HEADING.match(line.strip())
        if m and len(m.group(1)) == 1:
            return lines[:idx] + lines[idx + 1:], m.group(2).strip()
        break
    return lines, ""


def _style_table(table, header=True):
    """Shaded header, zebra body, hairline horizontals — and NO vertical rules.

    At 9.5pt a full grid reads as a cage: every cell gets four lines and the eye
    spends its effort on the ink rather than the values. A filled header row and
    horizontal hairlines carry the same structure with a fraction of it. Applied
    after the cells are filled, because `_add_inline` decides the runs.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    # With no header row, row 0 is data: the zebra shifts by one and the
    # first column carries the labels, in bold.
    first = 0 if header else -1
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if r != first and (r - first) % 2 == 0:
                _shade(cell, BAND_HEX)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    _face(run, BODY_FONT, 9.5,
                          color=((0xFF, 0xFF, 0xFF) if r == first else INK))
                    if r == first or (not header and c == 0):
                        run.font.bold = True

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        off = edge in ("left", "right", "insideV")
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none" if off else "single")
        el.set(qn("w:sz"), "0" if off else "6")
        el.set(qn("w:color"), "auto" if off else RULE_HEX)
        borders.append(el)
    table._tbl.tblPr.append(borders)

    # A table that runs onto the next page repeats its header row there, and no
    # row is cut in half across a page break: page 2 of a BBP showed a table
    # continuing with no idea which column was which (2026-09-28).
    for r, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if r == first:
            hdr = OxmlElement("w:tblHeader")
            tr_pr.append(hdr)
            # ...and the header row stays with the first data row: Word keeps a
            # row with the next through its paragraphs' keep-with-next, and
            # without it a page could end on a header row with nothing under it.
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True

    margins = OxmlElement("w:tblCellMar")
    for side, twips in (("top", 60), ("bottom", 60), ("left", 130), ("right", 130)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    table._tbl.tblPr.append(margins)


def _size_columns(table, headers, rows, measure):
    """Give each column the width its content asks for, not an equal share.

    Word's default split gave a `No` column holding "K1" the same width as the
    decision beside it, so the text column squeezed to a third of the page and a
    meeting note ran to six pages (2026-09-28). A column whose longest entry is
    short keeps its natural width; what is left goes to the text columns in
    proportion to how much they hold on average.
    """
    from docx.shared import Emu

    def plain(t):
        return re.sub(r"[*`_]|\[([^\]]*)\]\([^)]*\)", lambda m: m.group(1) or "", t or "")

    ncols = len(headers)
    char, pad = 72000, 180000       # EMU: ~0.20 cm per 9.5pt Segoe UI char, 0.5 cm cell padding
    cols = [[plain(headers[c])] + [plain(r[c]) if c < len(r) else "" for r in rows]
            for c in range(ncols)]
    longest = [max((len(v) for v in col), default=0) for col in cols]
    header_word = [max((len(w) for w in plain(headers[c]).split()), default=0) for c in range(ncols)]
    natural = [int(max(longest[c], header_word[c]) * char + pad) for c in range(ncols)]
    total = int(measure)
    if sum(natural) <= total:
        widths = natural[:]
        widths[longest.index(max(longest))] += total - sum(natural)
    else:
        # A column whose longest entry fits in ~22 characters (a code, a name,
        # "Ana veri") keeps a width that shows it on one line; only columns that
        # hold running text share what is left, by how much they hold on average.
        cap = 22
        floor = [int(min(max(longest[c], header_word[c]), cap) * char + pad) for c in range(ncols)]
        prose = [c for c in range(ncols) if longest[c] > cap]
        fixed = sum(floor[c] for c in range(ncols) if c not in prose)
        if not prose or fixed >= total * 0.7:
            return                                  # nothing sensible to do; keep Word's split
        avg = {c: max(1, sum(len(v) for v in cols[c][1:]) / max(1, len(cols[c]) - 1)) for c in prose}
        share = sum(avg.values())
        widths = [floor[c] if c not in prose else int((total - fixed) * avg[c] / share)
                  for c in range(ncols)]
    # Never wider than the measure: a table past the right margin is cut off in
    # Word and in print (a BBP table with four text columns did exactly that).
    # Squeeze the text columns first, down to their longest word; then everything.
    # No column narrower than its longest word: below that Word breaks words
    # letter by letter ("Sipari/ş"), which reads worse than a longer table.
    word = [int(min(max((len(w) for v in cols[c] for w in v.split()), default=1), 18) * char + pad)
            for c in range(ncols)]
    widths = [max(widths[c], word[c]) for c in range(ncols)]
    over = sum(widths) - total
    if over > 0:
        give = [max(0, widths[c] - word[c]) for c in range(ncols)]
        room = sum(give)
        if room:
            take = min(over, room)
            widths = [widths[c] - int(take * give[c] / room) for c in range(ncols)]
        if sum(widths) > total:
            k = total / sum(widths)
            widths = [int(w * k) for w in widths]
    table.autofit = False
    grid = table._tbl.tblGrid
    for c, gc in enumerate(grid.gridCol_lst):
        gc.w = Emu(widths[c])
    for row in table.rows:
        for c, cell in enumerate(row.cells):
            cell.width = Emu(widths[c])


def _keep_together(doc):
    """Nothing that introduces something is left at the foot of a page.

    A heading, a bold one-line lead-in ("**2.1.10 Açık Konular**") and any
    paragraph directly above a table keep with what follows. Without this a
    BBP ended pages on "2.1.10 Açık Konular" with its table on the next page.
    """
    from docx.text.paragraph import Paragraph

    body = [el for el in doc.element.body if el.tag.split("}")[-1] in ("p", "tbl")]
    for k, el in enumerate(body[:-1]):
        if not el.tag.endswith("}p"):
            continue
        p = Paragraph(el, doc)
        text = p.text.strip()
        if not text:
            continue
        heading = p.style is not None and p.style.name.startswith("Heading")
        runs = [r for r in p.runs if r.text.strip()]
        lead_in = (len(text) <= 120 and runs and all(r.bold for r in runs))
        before_table = body[k + 1].tag.endswith("}tbl")
        if heading or lead_in or before_table or text.endswith(":"):
            p.paragraph_format.keep_with_next = True


def _insert_toc(doc, anchor, heading_title):
    """A table of contents after the title block: a TOC field whose cached
    result already lists the headings as links, so it reads right even where
    Word never opens the file; `_word_page_numbers` then fills in page
    numbers on a machine that has Word."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph

    from toc import toc_depth
    # (heading, level, characters of text under it) for every h1-h3: how deep
    # the list goes depends on how much each deepest entry holds.
    sized, cur = [], None
    for el in doc.element.body:
        tag = el.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(el, anchor._parent)
            name = para.style.name if para.style is not None else ""
            if name in ("Heading 1", "Heading 2", "Heading 3") and para.text.strip():
                cur = [para, int(name[-1]), 0]
                sized.append(cur)
                continue
            if name.startswith("Heading"):
                continue
        if cur is not None and tag in ("p", "tbl"):
            cur[2] += len("".join(el.itertext()))
    depth = toc_depth([(lvl, chars) for _, lvl, chars in sized])
    heads = [para for para, lvl, _ in sized if lvl <= depth]
    if not heads:
        return 0
    marks = []
    for n, h in enumerate(heads, start=1):
        name = f"_Toc{100000 + n}"
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(n))
        start.set(qn("w:name"), name)
        end.set(qn("w:id"), str(n))
        h._p.insert(1 if h._p.pPr is not None else 0, start)
        h._p.append(end)
        marks.append((name, int(h.style.name[-1]), h.text.strip()))

    def fld(kind, instr=None):
        r = OxmlElement("w:r")
        if instr is None:
            c = OxmlElement("w:fldChar")
            c.set(qn("w:fldCharType"), kind)
            r.append(c)
        else:
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = instr
            r.append(t)
        return r

    # The title of the block, styled like a heading but not one: a TOC that
    # lists its own title reads as a mistake.
    title = OxmlElement("w:p")
    anchor._p.addnext(title)
    tp = Paragraph(title, anchor._parent)
    tp.paragraph_format.space_before = Pt(6)
    tp.paragraph_format.space_after = Pt(6)
    tp.paragraph_format.keep_with_next = True
    _face(tp.add_run(heading_title), _HEAD_FONT, 13, False, _ACCENT)

    last = title
    min_level = min(lvl for _, lvl, _ in marks)
    for k, (name, lvl, text) in enumerate(marks):
        el = OxmlElement("w:p")
        last.addnext(el)
        last = el
        para = Paragraph(el, anchor._parent)
        pf = para.paragraph_format
        pf.left_indent = Cm(0.6 * (lvl - min_level))
        pf.space_after = Pt(2)
        pf.line_spacing = 1.15
        pf.tab_stops.add_tab_stop(Cm(15.6), alignment=2, leader=1)   # right, dots (Word's update)
        if k == 0:
            el.append(fld("begin"))
            el.append(fld(None, rf'TOC \o "1-{depth}" \h \z \u'))
            el.append(fld("separate"))
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), name)
        link.set(qn("w:history"), "1")
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(a), BODY_FONT)
        rpr.append(rf)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "20" if lvl > min_level else "21")
        rpr.append(sz)
        if lvl == min_level:
            b = OxmlElement("w:b")
            rpr.append(b)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        run.append(t)
        link.append(run)
        el.append(link)
        if k == len(marks) - 1:
            el.append(fld("end"))
    spacer = OxmlElement("w:p")
    last.addnext(spacer)
    Paragraph(spacer, anchor._parent).paragraph_format.space_after = Pt(8)
    return len(marks)


def _word_page_numbers(path) -> dict:
    """Ask Word which page each contents entry's heading lands on -- read only.

    Only Word knows where a page breaks. It opens the file READ-ONLY, invisibly,
    reads the page of every `_Toc` bookmark, and closes without saving. Not
    saving is the point: an earlier version let Word update the contents and
    save, and on an NTT machine that applies the sensitivity label -- a green
    "RESTRICTED" in the header of the document (2026-09-28).

    Returns {bookmark: page}; {} off Windows, without Word, on any failure, or
    with NTT_DOCX_NO_WORD set -- the contents list then stays a linked list
    without numbers. Never fatal.
    """
    import json
    import subprocess
    if os.name != "nt" or os.environ.get("NTT_DOCX_NO_WORD"):
        return {}
    ps = ("$ErrorActionPreference='Stop';$w=New-Object -ComObject Word.Application;"
          "$w.Visible=$false;$w.DisplayAlerts=0;$out=@{};"
          "try{$d=$w.Documents.Open($env:NTT_DOCX_PATH,$false,$true,$false);"
          "$d.Bookmarks.ShowHidden=$true;$d.Repaginate();"
          "foreach($b in $d.Bookmarks){if($b.Name -like '_Toc1*'){"
          "$out[$b.Name]=[int]$b.Range.Information(3)}};"
          "$d.Close(0)}finally{$w.Quit()};"
          "[Console]::Out.Write(($out|ConvertTo-Json -Compress))")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                           env={**os.environ, "NTT_DOCX_PATH": str(Path(path).resolve())},
                           capture_output=True, timeout=180)
        if r.returncode != 0:
            return {}
        got = json.loads(r.stdout.decode("utf-8", "replace") or "{}")
        return {str(k): int(v) for k, v in got.items()} if isinstance(got, dict) else {}
    except (OSError, ValueError, subprocess.SubprocessError):
        return {}


def _ask_word_to_number(path) -> None:
    """No Word here to measure pages (a server, another OS): set the document's
    update-fields flag, so Word fills the contents numbers when the file is
    opened (it asks once; in Protected View after "Enable editing"). Never fatal."""
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document(str(path))
        st = doc.settings.element
        if st.find(qn("w:updateFields")) is None:
            el = OxmlElement("w:updateFields")
            el.set(qn("w:val"), "true")
            # schema order: updateFields sits before hdrShapeDefaults / footnotePr / endnotePr / compat
            after = next((c for c in st if c.tag.split("}")[-1] in
                          ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                           "mathPr", "themeFontLang", "clrSchemeMapping", "decimalSymbol", "listSeparator")), None)
            if after is not None:
                after.addprevious(el)
            else:
                st.append(el)
            doc.save(str(path))
    except Exception:
        pass


def _number_toc(path, pages: dict) -> None:
    """Write the page numbers into the contents entries: a right tab with a
    dotted leader, then the number -- the way a book sets it. Adding them does
    not move a page break, so the numbers Word measured stay true."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(path)
    for link in doc.element.body.iter(qn("w:hyperlink")):
        name = link.get(qn("w:anchor"))
        if not name or name not in pages:
            continue
        para = link.getparent()
        for text in (None, str(pages[name])):
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(a), BODY_FONT)
            rpr.append(rf)
            col = OxmlElement("w:color")
            col.set(qn("w:val"), "%02X%02X%02X" % MUTED)
            rpr.append(col)
            r.append(rpr)
            if text is None:
                r.append(OxmlElement("w:tab"))
            else:
                t = OxmlElement("w:t")
                t.text = text
                r.append(t)
            link.addnext(r)
            link = r
    doc.save(path)


MERMAID_FENCE = re.compile(r"^```\s*mermaid\s*$", re.IGNORECASE)


def _render_mermaid(blocks: list[str], out_dir: Path) -> list:
    """Each ```mermaid block -> a PNG, drawn by mermaid.js in Chromium.

    Word cannot draw a diagram from its source, and until 2026-09-28 the TS's
    seven diagrams reached the Word edition as grey code: the PDF showed the
    process flow, the Word file a listing of it. Same engine and colours as
    office-pdf, so both editions show the same picture. Returns one path per
    block, or None where it could not be drawn -- the caller then keeps the code
    block and says so, rather than dropping the diagram.
    """
    none = [None] * len(blocks)
    if not blocks or os.environ.get("NTT_DOCX_NO_BROWSER"):
        return none
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("[office-docx] WARNING: Playwright is not installed; mermaid diagrams "
              "stay as code. py -m pip install playwright && py -m playwright "
              "install chromium", file=sys.stderr)
        return none
    import html as _h
    from theme import DEEP_HEX, MUTED_HEX, TINT_HEX
    divs = "".join(f'<div class="d" id="d{k}"><pre class="mermaid">{_h.escape(b)}</pre></div>'
                   for k, b in enumerate(blocks))
    page_html = (
        "<!doctype html><html><head><meta charset='utf-8'><style>"
        "body{margin:0;background:#fff;font-family:'Segoe UI',sans-serif}"
        ".d{display:inline-block;padding:10px;background:#fff}"
        "pre.mermaid{margin:0}</style></head><body>" + divs +
        "<script type='module'>"
        "import mermaid from 'https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs';"
        "mermaid.initialize({startOnLoad:false,theme:'base',securityLevel:'loose',"
        f"themeVariables:{{primaryColor:'#{TINT_HEX}',primaryBorderColor:'#{ACCENT_HEX}',"
        f"primaryTextColor:'#{DEEP_HEX}',lineColor:'#{MUTED_HEX}',"
        "fontFamily:'Segoe UI, sans-serif',fontSize:'14px'}});"
        "window.__mmDone=false;(async()=>{try{await mermaid.run()}catch(e){}"
        # mermaid sizes an svg to its container (max-width:100%), and an
        # inline-block container is sized by its content: the two shrink each
        # other and the diagram came out a thumbnail with unreadable labels.
        # Give each svg its own drawn width; Word shrinks it to the column.
        "document.querySelectorAll('.d svg').forEach(s=>{const v=s.viewBox.baseVal;"
        "if(v&&v.width){s.style.maxWidth='none';s.setAttribute('width',v.width);"
        "s.setAttribute('height',v.height)}});"
        "window.__mmDone=true})();</script></body></html>")
    out = list(none)
    try:
        src = out_dir / "mermaid.html"
        src.write_text(page_html, encoding="utf-8")
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={"width": 1100, "height": 900},
                                    device_scale_factor=2)
            page.goto(src.resolve().as_uri())
            page.wait_for_function("window.__mmDone === true", timeout=60000)
            for k in range(len(blocks)):
                el = page.query_selector(f"#d{k}")
                # A block mermaid could not parse is left as its text, not an svg.
                if el is None or page.query_selector(f"#d{k} svg") is None:
                    continue
                png = out_dir / f"mermaid-{k}.png"
                el.screenshot(path=str(png))
                out[k] = png
            browser.close()
    except Exception as e:  # noqa: BLE001 -- no Chromium, offline CDN, timeout
        print(f"[office-docx] WARNING: mermaid diagrams stay as code ({e.__class__.__name__}: "
              f"{str(e).splitlines()[0][:120] if str(e) else ''})", file=sys.stderr)
    missing = sum(1 for x in out if x is None)
    if missing:
        print(f"[office-docx] WARNING: {missing} of {len(blocks)} mermaid diagram(s) "
              "could not be drawn and stay as code", file=sys.stderr)
    return out


def _build_header(section, measure, logo=True):
    """The NTT mark, top right, on every page — and nothing else up there.

    A running head carrying the document id beside a logo that already repeats
    is furniture, not information, so the id lives in the footer instead. The
    logo is a SECTION property: it repeats however long the document grows,
    without anything being pasted per page.

    Falls back to a text wordmark when the asset is MISSING — a document that
    loses its logo file should look plainer, never broken. `logo=False` is a
    different thing and gets a different answer: the caller is converting a
    document that is not ours, so the header stays empty rather than printing
    the brand in letters instead of in pixels.
    """
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm, Emu

    if not logo:
        return
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.paragraph_format.space_after = 0
    # Right-aligned paragraph, not a tab to a right stop: Word's Header style
    # carries its own centre tab stop, so the tab landed there and the mark sat
    # centred on every document until 2026-09-28.
    p.alignment = 2                                   # right
    mark = logo_path()
    if mark:
        p.add_run().add_picture(str(mark), width=Cm(3.1))
    else:
        _face(p.add_run("NTT DATA"), BODY_FONT, 12, True, ACCENT)   # the MARK, never recoloured


def _build_footer(section, measure, left="", classification="", page_word="Sayfa"):
    """What a page torn out of the printed copy still has to say: which document
    it belongs to, how far it may travel, and where it sits in the whole."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Emu, Pt

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    # Word's Footer style brings its own centre and right stops (8.25 / 16.5 cm,
    # for Letter). Left in place they catch our tabs first, and the page number
    # sat mid-page on every document until 2026-09-28. Clear them; ours rule.
    try:
        p.style.paragraph_format.tab_stops.clear_all()
    except Exception:
        pass
    stops = p.paragraph_format.tab_stops
    stops.add_tab_stop(Emu(int(measure) // 2), WD_TAB_ALIGNMENT.CENTER)
    stops.add_tab_stop(Emu(int(measure)), WD_TAB_ALIGNMENT.RIGHT)
    # The title shares the line with the classification at the centre stop; a
    # long one ran into it ("...RaporuGizli"). Cut it, as the PDF footer does.
    if len(left) > 40:
        left = left[:38].rstrip(" —-·") + "…"
    _face(p.add_run(left), BODY_FONT, 8, False, MUTED)
    p.add_run("\t")
    _face(p.add_run(classification), BODY_FONT, 8, False, MUTED)
    p.add_run("\t")
    _face(p.add_run(page_word + " "), BODY_FONT, 8, False, MUTED)
    _face(_field(p, "PAGE"), BODY_FONT, 8, False, MUTED)
    _face(p.add_run(" / "), BODY_FONT, 8, False, MUTED)
    _face(_field(p, "NUMPAGES"), BODY_FONT, 8, False, MUTED)
    _rule(p, "top", RULE_HEX, size=6, space=6)


def build(md_text, out_path, title, base_dir=".", eyebrow="",
          classification="", logo=True, accent="", heading_font="",
          toc="auto", toc_title="İçindekiler", page_word="Sayfa"):
    from docx import Document
    from docx.shared import Cm, Emu, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    _apply_palette(accent, heading_font)
    base = Path(base_dir)
    # No --title: the document's own first `# ...` line is the title, as in
    # office-pdf. Without this a document rendered with only --eyebrow got no
    # title block, its `#` printed as a heading and the footer had no name.
    if not title:
        for ln in strip_html_comments(md_text).splitlines():
            if ln.strip():
                m = HEADING.match(ln.strip())
                if m and len(m.group(1)) == 1:
                    title = _unpark(re.sub(r"[*`]", "", _park(m.group(2)))).strip()
                break

    doc = Document()

    # python-docx ships a US Letter template with 3.17 cm margins. Everything
    # this kit produces is read and printed on A4, so set the page rather than
    # inherit it: on Letter the text reflows when someone prints A4, and the old
    # hardcoded 15.5 cm image cap was WIDER than the real 15.24 cm column, so
    # oversized screenshots ran into the margin.
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.7)
        section.top_margin, section.bottom_margin = Cm(2.6), Cm(2.2)
        section.header_distance = section.footer_distance = Cm(1.1)
    first = doc.sections[0]
    max_width = Emu(int(first.page_width)
                    - int(first.left_margin) - int(first.right_margin))

    _build_header(first, max_width, logo=logo)
    _build_footer(first, max_width, left=title, classification=classification,
                  page_word=page_word)

    # Airiness is a set of numbers, not a mood. 1.35 leading (was 1.15) and a
    # ~16 cm measure -- about 80 characters -- is the width the eye tracks
    # without losing its place on the way back to the left margin.
    normal = doc.styles["Normal"]
    _face(normal, BODY_FONT, 10.5, color=INK)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(10)

    # Headings in the accent, sized so three levels are told apart by weight and
    # space rather than by colour alone, and given room to separate sections
    # instead of sitting on the paragraph above.
    # Four to six keep stepping down, in size and then in colour. Each is set
    # explicitly: Word's own Heading 4 is italic and 4-6 follow the template's
    # BLUE theme accent, a second colour system inside the document.
    # Each level must look smaller than the one above it: at 11pt navy
    # Semibold, Heading 3 ("2.1 Current Situation") read lighter than the bold
    # 10.5pt line under it (2026-09-28). So 3 is larger and stays navy, and 4
    # onward step down to the text ink, then to muted.
    levels = ((16, 22), (13.5, 18), (12, 14), (11, 12), (10.5, 8), (10, 8))
    for lvl, (size, before) in enumerate(levels, start=1):
        try:
            style = doc.styles[f"Heading {lvl}"]
            _face(style, _HEAD_FONT, size,
                  color=(_ACCENT if lvl <= 2 else _DEEP if lvl == 3
                         else INK if lvl <= 5 else MUTED))
            style.font.bold = False          # Semibold IS the weight; bold doubles it
            style.font.italic = False
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(8)
            style.paragraph_format.line_spacing = 1.0
        except Exception:
            pass

    if title:
        if eyebrow:
            eb = doc.add_paragraph()
            eb.paragraph_format.space_after = Pt(3)
            eb.paragraph_format.line_spacing = 1.0
            _face(eb.add_run(eyebrow), BODY_FONT, 8, True, _ACCENT).font.all_caps = True
        # Not add_heading(level=0): Word's Title style carries its own font and
        # its own bottom border, both of which would fight this theme.
        h = doc.add_paragraph()
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing = 1.0
        # The title takes the heading face too. A page whose H1 is Cambria and
        # whose title is Segoe UI Light reads as two documents stapled together.
        _face(h.add_run(title), LIGHT_FONT if _HEAD_FONT == SEMI_FONT else _HEAD_FONT,
              22, False, _DEEP)
        _rule(h, "bottom", _ACCENT_HEX, size=12, space=9)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(10)
    toc_anchor = doc.paragraphs[-1] if doc.paragraphs else None

    lines, taken = take_leading_title(strip_html_comments(md_text).splitlines(),
                                      title)
    lines = _park_outside_fences(lines)
    if taken and _plain(taken) != _plain(title):
        # Not an error: an FS opens with `# FS-FI-012 — <title>` and passes the
        # business title alone. Said out loud so nobody wonders where it went.
        print(f"[office-docx] leading '# {taken}' replaced by --title '{title}'",
              file=sys.stderr)
    # Diagrams are drawn up front, in one browser session for all of them.
    import tempfile
    mm_tmp = tempfile.TemporaryDirectory()
    mm_src, k = [], 0
    while k < len(lines):
        if MERMAID_FENCE.match(lines[k].strip()):
            body, k = [], k + 1
            while k < len(lines) and not lines[k].strip().startswith("```"):
                body.append(lines[k])
                k += 1
            mm_src.append("\n".join(body))
        k += 1
    mm_png = iter(_render_mermaid(mm_src, Path(mm_tmp.name)))
    # A diagram never takes more than this much of a page's height.
    max_height = Cm(20)

    i, n = 0, len(lines)
    # Numbering instance of the ordered list currently being emitted. Only a
    # structural block (heading, table, code, page break) starts a new list;
    # images, notes and quotes between two steps must not restart the count.
    num_id = None
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # blank
        if not stripped:
            i += 1
            continue

        # horizontal rule -> a hairline, as the PDF draws it. It was a page
        # break until 2026-09-29: an As-Built with a rule between chapters ran
        # to 22 pages in Word against 12 in PDF, most of them half empty, and a
        # break falling on the last line of a page left a blank page behind it.
        # No skill used it to mean "new page".
        if stripped in ("---", "***", "___"):
            num_id = None
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = Pt(2)
            _rule(p, "bottom", RULE_HEX, size=6, space=1)
            i += 1
            continue

        # fenced code block
        if stripped.startswith("```"):
            is_mermaid = bool(MERMAID_FENCE.match(stripped))
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # closing fence
            num_id = None
            png = next(mm_png, None) if is_mermaid else None
            if png is not None:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(12)
                p.add_run().add_picture(str(png))
                shape = doc.inline_shapes[-1]
                # Drawn at 2x for print sharpness, so its natural size is half its
                # pixels at 96 per inch (9525 EMU a CSS pixel) -- not whatever DPI
                # python-docx assumes for a PNG that states none.
                import struct
                pw, ph = struct.unpack(">II", png.read_bytes()[16:24])
                w, h = pw * 9525 // 2, ph * 9525 // 2
                ratio = min(1.0, int(max_width) / w, int(max_height) / h)
                shape.width, shape.height = int(w * ratio), int(h * ratio)
                continue
            p = doc.add_paragraph()
            # Set apart from the prose: indented, single-spaced (code lines are
            # meant to be read as a block, not as paragraphs), with air around it.
            pf = p.paragraph_format
            pf.left_indent = Cm(0.45)
            pf.line_spacing = 1.15
            pf.space_before = Pt(8)
            pf.space_after = Pt(12)
            _face(p.add_run("\n".join(code)), MONO_FONT, 9, color=INK)
            _shade(p, BAND_HEX)
            continue

        # image line: one ![alt](path), or several separated by spaces, which
        # render as a ROW. Markdown already treats images as inline, so the
        # source reads the way the page looks: three logos on one line are three
        # logos side by side, and nothing new has to be learned to get that.
        # A path may be `<in angle brackets>` or percent-encoded (%20): project
        # folders have spaces ("04 Ekran Taslagi"), and a bare space ends a
        # Markdown link -- the FS's mock-up pictures were dropped that way.
        img = r"!\[([^\]]*)\]\((<[^>]+>|[^)\s]+)\)"
        shots = re.findall(img, stripped)
        if shots and re.fullmatch(rf"(?:{img}\s*)+", stripped):
            from urllib.parse import unquote
            paths = []
            for _alt, ref in shots:
                src = Path(unquote(ref.strip("<>")))
                if not src.is_absolute():
                    src = base / src
                if src.is_file():
                    paths.append(src)
                else:
                    print(f"[office-docx] WARNING: image not found: {src}",
                          file=sys.stderr)
            if paths:
                gap = Pt(6)
                # Share the column between them, minus the gaps. Each picture is
                # still only ever shrunk: a small logo keeps its own size rather
                # than being blown up to fill a third of the page.
                budget = Emu(int(max_width) - int(gap) * (len(paths) - 1))
                share = Emu(int(budget) // len(paths))
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # NOT `n`: that is the main loop's line count, and shadowing it
                # ended the document at the first image row.
                for idx, src in enumerate(paths):
                    if idx:
                        p.add_run(" ")
                    run = p.add_run()
                    run.add_picture(str(src))
                    shape = doc.inline_shapes[-1]
                    if shape.width > share:
                        ratio = int(share) / int(shape.width)
                        shape.width = int(int(shape.width) * ratio)
                        shape.height = int(int(shape.height) * ratio)
            i += 1
            continue

        # heading
        m = HEADING.match(stripped)
        if m:
            num_id = None
            # Build the heading empty and fill it through the inline renderer:
            # add_heading(text) writes the string verbatim, so a heading naming a
            # file or a flag came out with its backticks showing.
            h = doc.add_heading("", level=len(m.group(1)))
            _add_inline(h, m.group(2).strip())
            i += 1
            continue

        # blockquote -> callout: a bar in the accent, indented, quieter text.
        # Word's "Intense Quote" was centred, italic and followed the template's
        # BLUE theme accent, so a document had two unrelated colour systems in it.
        # A quote hard-wrapped over several `>` lines is ONE paragraph; a bare
        # `>` line separates paragraphs inside it. Each line used to become its
        # own paragraph, so a four-line note printed as four spaced fragments.
        if stripped.startswith(">"):
            paras, cur = [], []
            while i < n and lines[i].strip().startswith(">"):
                t = lines[i].strip()[1:].strip()
                if t:
                    cur.append(t)
                elif cur:
                    paras.append(cur)
                    cur = []
                i += 1
            if cur:
                paras.append(cur)
            for k, chunk in enumerate(paras):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.55)
                p.paragraph_format.space_after = Pt(12 if k == len(paras) - 1 else 4)
                _add_inline(p, " ".join(chunk))
                for run in p.runs:
                    _face(run, BODY_FONT, 10, color=MUTED)
                _rule(p, "left", _ACCENT_HEX, size=18, space=11)
            continue

        # GFM table: header row + separator row
        if "|" in line and i + 1 < n and _is_table_sep(lines[i + 1]):
            num_id = None
            headers = _split_row(line)
            i += 2
            rows = []
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(_split_row(lines[i]))
                i += 1
            # `| | |` over a key/value block is Markdown's way of writing a
            # table with no header; it printed as an empty blue bar.
            headless = not any(h.strip() for h in headers)
            table = doc.add_table(rows=0 if headless else 1, cols=len(headers))
            table.style = "Table Grid"
            for c, htext in enumerate([] if headless else headers):
                cell = table.rows[0].cells[c]
                cell.text = ""
                _add_inline(cell.paragraphs[0], htext)
                _shade(cell, _ACCENT_HEX)
            for row in rows:
                cells = table.add_row().cells
                for c in range(len(headers)):
                    cells[c].text = ""
                    _add_inline(cells[c].paragraphs[0],
                                row[c] if c < len(row) else "")
            _style_table(table, header=not headless)
            _size_columns(table, headers, rows, max_width)
            continue

        # bullet / ordered list
        mb = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if mb:
            indent = len(mb.group(1)) // 2
            ordered = mb.group(2).endswith(".")
            style = "List Number" if ordered else "List Bullet"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.left_indent = Pt(18 * (indent + 1))
            if ordered:
                if num_id is None:          # first item of a new ordered list
                    num_id = _fresh_num_id(doc)
                _apply_num(p, num_id, indent)
            else:
                num_id = None
            # The item's own continuation lines -- a wrapped item in the source
            # -- belong to it, not to a numberless paragraph under it.
            item = [mb.group(3).strip()]
            i += 1
            while i < n and lines[i].strip() \
                    and not _BLOCK_START.match(lines[i].strip()) \
                    and not lines[i].lstrip().startswith("|"):
                item.append(lines[i].strip())
                i += 1
            _add_inline(p, " ".join(item))
            continue

        # plain paragraph (gather consecutive non-special lines)
        para = [stripped]
        i += 1
        while i < n and lines[i].strip() \
                and not _BLOCK_START.match(lines[i].strip()) \
                and "|" not in lines[i]:
            para.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        _add_inline(p, " ".join(para))

    # Stamp the source so a later run can tell whether this file is still the
    # markdown it was built from. A .docx is a build output that nothing else
    # links back to its source: edit the guide, forget to rebuild, and the Word
    # copy people download drifts with no way to notice. Lives in core
    # properties, which survive render_manual()'s hyperlink rewrite (that
    # touches word/_rels only).
    doc.core_properties.comments = f"{SOURCE_STAMP}{source_sha(md_text)}"

    _keep_together(doc)
    # A contents list pays its way in a long document and is clutter in a short
    # one: `auto` adds it from eight headings up.
    n_heads = sum(1 for p in doc.paragraphs if p.style is not None
                  and p.style.name in ("Heading 1", "Heading 2", "Heading 3"))
    want_toc = toc == "on" or (toc == "auto" and n_heads >= 8)
    if want_toc and toc_anchor is not None:
        _insert_toc(doc, toc_anchor, toc_title)

    # Word needs a paragraph after a closing table and adds a full-height empty
    # one when the file has none; when the table ends near the foot of a page
    # that paragraph alone makes a blank last page (a meeting note, 2026-09-28).
    # End on a 1pt paragraph with no spacing instead.
    body = doc.element.body
    blocks = [el for el in body if el.tag.split("}")[-1] in ("p", "tbl")]
    if blocks and blocks[-1].tag.endswith("}tbl"):
        from docx.shared import Pt
        tail = doc.add_paragraph()
        tail.paragraph_format.space_before = Pt(0)
        tail.paragraph_format.space_after = Pt(0)
        tail.paragraph_format.line_spacing = Pt(1)
        tail.add_run("").font.size = Pt(1)
    doc.save(out_path)
    if want_toc and toc_anchor is not None:
        pages = _word_page_numbers(out_path)
        if pages:
            _number_toc(out_path, pages)
        else:
            _ask_word_to_number(out_path)


# The fixed words a renderer prints itself, per document language (--lang). The master is
# lib/doc_language.json ("render"); scripts/test_doc_language.py holds this table equal to it.
LANG_WORDS = {"tr": {"toc_title": "İçindekiler", "page_word": "Sayfa"},
              "en": {"toc_title": "Contents", "page_word": "Page"}}


def resolve_lang_words(lang, toc_title, page_word):
    """--toc-title / --page-word win; otherwise --lang's words; otherwise Turkish, as before."""
    words = LANG_WORDS[lang or "tr"]
    return (toc_title if toc_title is not None else words["toc_title"],
            page_word if page_word is not None else words["page_word"])


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Markdown subset -> native .docx.")
    ap.add_argument("--input", required=True, help="source .md file")
    ap.add_argument("--output", required=True, help="target .docx")
    ap.add_argument("--title", default="",
                    help="document title; replaces a leading '# ...' line")
    ap.add_argument("--eyebrow", default="",
                    help="small caps line above the title, e.g. TEKNIK SPESIFIKASYON")
    ap.add_argument("--classification", default="",
                    help="footer centre, e.g. 'Confidential - Customer Restricted'")
    ap.add_argument("--no-logo", action="store_true",
                    help="omit the NTT mark (converting a document that is not ours)")
    ap.add_argument("--accent", default="",
                    help="override the house accent, e.g. 2E7D32; the deeper "
                         "title shade is derived from it")
    ap.add_argument("--heading-font", default="",
                    help="override the heading face, e.g. Cambria (for a surface "
                         "whose HTML edition is set in it)")
    ap.add_argument("--redact-pii", action="store_true",
                    help="mask TCKN / 10-11 digit tax IDs before writing")
    ap.add_argument("--toc", choices=("auto", "on", "off"), default="auto",
                    help="table of contents: auto = from eight headings up")
    ap.add_argument("--lang", choices=("tr", "en"), default=None,
                    help="document language: sets the contents heading and the footer's page "
                         "label (tr: İçindekiler / Sayfa, en: Contents / Page). Default tr.")
    ap.add_argument("--page-word", default=None,
                    help="footer page label; overrides --lang")
    ap.add_argument("--toc-title", default=None,
                    help="its heading; overrides --lang")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args(argv)
    args.toc_title, args.page_word = resolve_lang_words(args.lang, args.toc_title, args.page_word)

    _require_docx()
    if not os.path.isfile(args.input):
        sys.exit(f"ERROR: input not found: {args.input}")
    if os.path.exists(args.output) and not args.force:
        sys.exit(f"ERROR: {args.output} exists (use --force)")

    # utf-8-sig: PowerShell's Set-Content/Out-File prepend a BOM, which would
    # otherwise ride into the first heading as an invisible character.
    md_text = open(args.input, encoding="utf-8-sig").read()
    if args.redact_pii:
        if redact_text is None:
            sys.exit("ERROR: --redact-pii requested but lib/redact.py not importable")
        n = count_matches(md_text)
        md_text = redact_text(md_text)
        sys.stderr.write(f"[redact] masked {n} ID(s) before writing\n")

    build(md_text, args.output, args.title,
          base_dir=os.path.dirname(os.path.abspath(args.input)),
          eyebrow=args.eyebrow, classification=args.classification,
          logo=not args.no_logo, accent=args.accent,
          heading_font=args.heading_font, toc=args.toc, toc_title=args.toc_title,
          page_word=args.page_word)
    print(f"[office-docx] wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
