"""The NTT look for every .xlsx the kit writes -- one definition, like theme.py.

WHY THIS FILE EXISTS

Until 2026-09-28 no skill owned its workbook's look. fast-scan told the agent
"use office-excel-write or office-excel-report", so each run picked its own
colours and settings. A sample run came out in a navy that is not the house blue
(#003087 against theme.py's #0072BC), in Calibri while every .docx/.pdf is Segoe
UI, with filter arrows covering the header text, the question in the seventh of
nineteen columns, no column for the customer's answer, and a print setup that put
six columns on page one and none of them the question. None of that is a bug in
one file; it is what "each run decides" produces.

What a workbook built here always has:
  - a header block: the NTT mark, the document title, customer / project /
    version / date, and one line saying how to use the sheet;
  - a real Excel table (filter and sort on every column, grows with new rows),
    with the house blue header, zebra bands and hairlines from theme.py;
  - Segoe UI throughout, wrapped text aligned to the top, a little indent;
  - header cells padded so the filter arrow never covers the label;
  - columns the reader fills in (`input=True`) marked by colour and legend,
    with drop-down lists where the answers are a fixed set;
  - frozen header (and first column), no gridlines;
  - print: A4 landscape, fit to one page WIDE, header row repeated on every
    page, page numbers and the document name in the footer.

openpyxl only (office-tools already requires it). The logo needs Pillow, which
python-pptx brings; without it the mark is set in type, never missing.
"""
from __future__ import annotations

import datetime as _dt
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    from theme import (ACCENT_HEX, BODY_FONT, DEEP_HEX, INK_HEX, MUTED_HEX,
                       RULE_HEX, SEMI_FONT, TINT_HEX, logo_path)
except ImportError:                                   # imported as office-tools.lib
    from .theme import (ACCENT_HEX, BODY_FONT, DEEP_HEX, INK_HEX,  # type: ignore
                        MUTED_HEX, RULE_HEX, SEMI_FONT, TINT_HEX, logo_path)

# A cell the reader fills in. Warm, so it reads as "yours" against the blue
# system; the one colour here that is not on the blue hue, on purpose.
INPUT_HEX = "FFF8E1"
WHITE = "FFFFFF"

HEADER_ROW = 6           # rows 1-4 header block, 5 a spacer
FILTER_PAD = 3           # characters kept free for the filter arrow

_hair = Side(style="thin", color=RULE_HEX)
_BORDER = Border(bottom=_hair, top=_hair, left=_hair, right=_hair)


def _font(size=10, bold=False, color=INK_HEX, italic=False, semi=False):
    return Font(name=SEMI_FONT if semi else BODY_FONT, size=size, bold=bold,
                color=color, italic=italic)


def new_workbook(meta: dict) -> Workbook:
    """meta: title, customer, project, version, author (all optional)."""
    wb = Workbook()
    wb.remove(wb.active)
    p = wb.properties
    p.title = meta.get("title") or ""
    p.creator = meta.get("author") or "NTT DATA Business Solutions"
    p.company = "NTT DATA Business Solutions"
    p.subject = meta.get("project") or ""
    return wb


def _header_block(ws, meta: dict, ncols: int, legend: str) -> None:
    last = get_column_letter(max(ncols, 1))
    ws.row_dimensions[1].height = 34
    logo = logo_path()
    placed = False
    if logo:
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(str(logo))
            ratio = img.width / img.height
            img.height = 34
            img.width = int(34 * ratio)
            ws.add_image(img, "A1")
            placed = True
        except Exception:                         # no Pillow: set it in type
            placed = False
    if not placed:
        ws["A1"] = "NTT DATA"
        ws["A1"].font = Font(name=SEMI_FONT, size=14, bold=True, color=ACCENT_HEX)

    ws["A2"] = meta.get("title") or ""
    ws["A2"].font = Font(name=SEMI_FONT, size=16, color=DEEP_HEX)
    ws.row_dimensions[2].height = 26
    parts = []
    for label, key in (("Müşteri", "customer"), ("Proje", "project"),
                       ("Sürüm", "version"), ("Tarih", "date")):
        v = meta.get(key)
        if key == "date" and not v:
            v = _dt.date.today().strftime("%d.%m.%Y")
        if v:
            parts.append(f"{meta.get('labels', {}).get(key, label)}: {v}")
    ws["A3"] = "   ·   ".join(parts)
    ws["A3"].font = _font(10, color=MUTED_HEX)
    ws["A4"] = legend
    ws["A4"].font = _font(9, italic=True, color=MUTED_HEX)
    ws.row_dimensions[5].height = 6
    for r in (2, 3, 4):
        ws.merge_cells(f"A{r}:{last}{r}")
        ws[f"A{r}"].alignment = Alignment(vertical="center")


def place_mark(ws, anchor: str = "A1", height: int = 34) -> bool:
    """The NTT mark as a picture at `anchor`; False when the asset or Pillow is missing."""
    logo = logo_path()
    if not logo:
        return False
    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(str(logo))
        ratio = img.width / img.height
        img.height = height
        img.width = int(height * ratio)
        ws.add_image(img, anchor)
        return True
    except Exception:
        return False


def apply_house_style(ws, *, header_row: int = 1, ncols: int, last_row: int, title: str = "",
                      keep_fill: tuple = (), center: tuple = ()) -> None:
    """The house look on a table another writer already laid out: blue header, Segoe UI,
    hairlines, tinted zebra, A4 landscape on one page wide with the header repeated, and the
    title and page numbers in the footer. For a workbook whose layout its readers depend on
    (Kokpit's library and UAT results, 2026-09-29), where add_table_sheet's header block
    would move the table. Columns in keep_fill (1-based) hold fills of their own -- a status
    colour -- and are left out of the zebra."""
    for j in range(1, ncols + 1):
        c = ws.cell(header_row, j)
        c.font = Font(name=SEMI_FONT, size=10, color=WHITE)
        c.fill = PatternFill("solid", fgColor=ACCENT_HEX)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    first = header_row + 1
    for r in range(first, last_row + 1):
        for j in range(1, ncols + 1):
            c = ws.cell(r, j)
            f = c.font
            c.font = Font(name=BODY_FONT, size=10, bold=f.bold, italic=f.italic,
                          color=f.color if f.color is not None and f.color.rgb not in (None, "FF000000") else INK_HEX)
            c.border = Border(bottom=_hair)
            c.alignment = Alignment(wrap_text=True, vertical="top",
                                    horizontal="center" if j in center else None,
                                    indent=0 if j in center else 1)
    spans, start = [], None
    for j in range(1, ncols + 2):
        plain = j <= ncols and j not in keep_fill
        if plain and start is None:
            start = j
        elif not plain and start is not None:
            spans.append(f"{get_column_letter(start)}{first}:{get_column_letter(j - 1)}{last_row + 500}")
            start = None
    if spans and last_row >= first:
        ws.conditional_formatting.add(
            " ".join(spans),
            FormulaRule(formula=[f"AND(MOD(ROW()-{first},2)=1,$A{first}<>\"\")"],
                        fill=PatternFill("solid", fgColor=TINT_HEX, bgColor=TINT_HEX), stopIfTrue=False))
    ws.sheet_view.showGridLines = False
    _page_setup(ws, title, f"{header_row}:{header_row}")


def _page_setup(ws, title: str, repeat_rows: str | None = None) -> None:
    if repeat_rows:
        ws.print_title_rows = repeat_rows
    ps = ws.page_setup
    ps.orientation = "landscape"
    ps.paperSize = ws.PAPERSIZE_A4
    ps.fitToWidth = 1
    ps.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = (title or "").replace("&", "&&")
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "&P / &N"
    ws.oddFooter.right.size = 8


def add_table_sheet(wb: Workbook, name: str, columns: list[dict], rows: list[dict], *,
                    meta: dict, legend: str = "", freeze_cols: int = 1,
                    table_name: str | None = None, hide_empty: bool = False):
    """One sheet: header block + an Excel table.

    columns: dicts with
      key       field name in each row
      header    label shown
      width     column width in characters (default 16)
      input     True for a column the reader fills in
      choices   list of allowed values -> a drop-down (with input or controlled data)
      wrap      default True
      hidden    hide the column
    """
    ws = wb.create_sheet(title=name[:31])
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ncols = len(columns)
    _header_block(ws, meta, ncols, legend)

    h = HEADER_ROW
    for j, col in enumerate(columns, start=1):
        c = ws.cell(h, j, col["header"])
        c.font = Font(name=SEMI_FONT, size=10, color=WHITE)
        c.fill = PatternFill("solid", fgColor=DEEP_HEX if col.get("input") else ACCENT_HEX)
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        c.border = _BORDER
        width = max(col.get("width", 16), min(len(col["header"]), 28) + FILTER_PAD)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[h].height = 34

    # The title sits in a cell merged across the table, and a merged cell does
    # not spill: a title longer than a narrow table was cut mid-word (a
    # seven-column weekly report, 2026-09-29). Wrap it onto as many lines as it
    # needs; a 16pt character is about 1.6 column-width units.
    span = sum(ws.column_dimensions[get_column_letter(j)].width or 13
               for j in range(1, ncols + 1))
    need = len(meta.get("title") or "") * 1.6
    if need > span:
        lines = int(need // span) + 1
        ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 24 * lines

    first = h + 1
    for i, row in enumerate(rows):
        r = first + i
        for j, col in enumerate(columns, start=1):
            v = row.get(col["key"], "")
            c = ws.cell(r, j, "" if v is None else v)
            c.font = _font(10)
            c.alignment = Alignment(wrap_text=col.get("wrap", True), vertical="top", indent=1)
            c.border = _BORDER
            if col.get("input"):
                c.fill = PatternFill("solid", fgColor=INPUT_HEX)
    last_row = first + max(len(rows), 1) - 1
    last_col = get_column_letter(ncols)
    ref = f"A{h}:{last_col}{last_row}"

    tname = re.sub(r"[^A-Za-z0-9_]", "_", table_name or name)
    if not re.match(r"[A-Za-z_]", tname):
        tname = "T_" + tname
    tab = Table(displayName=tname[:250], ref=ref)
    # Stripes and header colour are ours (theme.py); the table brings filter,
    # sort and growth. TableStyleLight1 is the neutral one, so nothing of the
    # Office theme's own blue shows through.
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False,
                                        showColumnStripes=False, showFirstColumn=False,
                                        showLastColumn=False)
    ws.add_table(tab)

    # Zebra by formula, over room to grow, so a row added by hand is banded too.
    # TINT, not theme.py's BAND: BAND is set for a lit Word page and vanishes on
    # a spreadsheet (measured on the first Fast Scan render, 2026-09-28).
    band_to = last_row + 500
    # Input columns keep their own colour: a band painted over them would hide
    # every other "fill this in" cell.
    spans, start = [], None
    for j, col in enumerate(columns + [{"input": True}], start=1):
        if not col.get("input") and start is None:
            start = j
        elif col.get("input") and start is not None:
            spans.append(f"{get_column_letter(start)}{first}:{get_column_letter(j - 1)}{band_to}")
            start = None
    ws.conditional_formatting.add(
        " ".join(spans) or f"A{first}:{last_col}{band_to}",
        FormulaRule(formula=[f"AND(MOD(ROW()-{first},2)=1,$A{first}<>\"\")"],
                    # A conditional fill is painted from bgColor, not fgColor (Excel's
                    # rule for differential formats); set both or it stays white.
                    fill=PatternFill("solid", fgColor=TINT_HEX, bgColor=TINT_HEX),
                    stopIfTrue=False))

    for j, col in enumerate(columns, start=1):
        letter = get_column_letter(j)
        if col.get("choices"):
            opts = [str(o).replace(",", " ") for o in col["choices"] if str(o).strip()]
            if opts:
                dv = DataValidation(type="list", formula1='"' + ",".join(opts)[:250] + '"',
                                    allow_blank=True, showDropDown=False)
                dv.error = "Listeden bir değer seçin."
                dv.errorStyle = "warning"
                ws.add_data_validation(dv)
                dv.add(f"{letter}{first}:{letter}{band_to}")
        empty = hide_empty and all(str(r.get(col["key"], "") or "").strip() == "" for r in rows)
        if col.get("hidden") or (empty and not col.get("input")):
            ws.column_dimensions[letter].hidden = True

    ws.freeze_panes = ws.cell(first, freeze_cols + 1)
    ws.print_title_rows = f"{h}:{h}"
    ws.print_area = f"A1:{last_col}{last_row}"
    ps = ws.page_setup
    ps.orientation = "landscape"
    ps.paperSize = ws.PAPERSIZE_A4
    ps.fitToWidth = 1
    ps.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    doc = (meta.get("title") or "").replace("&", "&&")
    ws.oddFooter.left.text = doc
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "&P / &N"
    ws.oddFooter.right.size = 8
    return ws
