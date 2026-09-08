"""SAP Readiness Check XLSX per-tile → list[ReadinessTile].

RC XLSX convention (per research-readiness-check-verified.md):
- Rows 1-4: merged title cells (SAP banner + metadata)
- Row 5: column headers
- Row 6+: data
- Filename pattern: tile-<tile-name>.xlsx
"""

from pathlib import Path
from typing import cast

from openpyxl import load_workbook

from scripts.schemas import ReadinessTile


class ReadinessIngester:
    def load_dir(self, dir_path: Path) -> list[ReadinessTile]:
        if not dir_path.is_dir():
            return []
        tiles: list[ReadinessTile] = []
        for xlsx in sorted(dir_path.glob("tile-*.xlsx")):
            tile_name = xlsx.stem.removeprefix("tile-")
            tiles.append(ReadinessTile(tile_name=tile_name, rows=self._parse(xlsx)))
        return tiles

    @staticmethod
    def _parse(xlsx: Path) -> list[dict[str, str | int | float | None]]:
        wb = load_workbook(xlsx, read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return []
        rows: list[dict[str, str | int | float | None]] = []
        headers: list[str] = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == 5:
                headers = [str(c) for c in row if c is not None]
            elif i > 5:
                if all(c is None for c in row):
                    continue
                rows.append({
                    h: cast(str | int | float | None, row[idx]) if idx < len(row) else None
                    for idx, h in enumerate(headers)
                })
        return rows
